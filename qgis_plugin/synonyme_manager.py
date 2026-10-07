# -*- coding: utf-8 -*-
"""
Synonyme verwalten – Kernlogik (ohne UI).
=========================================

Pflege der Tabelle ``Arten_Synonyme`` (Referenzlisten).

Der schwierige Teil beim Ergänzen eines Synonyms ist nicht das Speichern,
sondern die korrekte Auflösung:

* ``akzeptiert_entityid`` muss auf eine real existierende Art zeigen,
* ``parentid`` muss eine **gültige** Artengruppe sein (``Artengruppen.listitemid``) –
  sonst ist der Eintrag über die Kaskade im Formular nie auffindbar.

Genau hier lagen die historischen Fehler (verwaiste Gruppen-IDs, Zielarten die
nicht in der Liste stehen). Dieses Modul erzwingt beides.

Alle Zugriffe laufen über die QGIS-Layer-API, damit GeoPackage **und** PostGIS
gleichermaßen unterstützt werden.
"""
import csv
import os
import re
from datetime import date

from qgis.PyQt.QtCore import QVariant
from qgis.core import (QgsApplication, QgsProject, QgsVectorLayer, QgsFeature,
                       QgsField, QgsFeatureRequest, QgsMessageLog)

TAG = "FT-Synonyme"

REF_LAYER = "Arten"
SYN_LAYER = "Arten_Synonyme"
AG_LAYER = "Artengruppen"


def log(msg):
    try:
        QgsMessageLog.logMessage(str(msg), TAG)
    except Exception:
        pass


def normalize(text):
    """Kleinschreibung + Umlautfaltung + Leerraum normalisieren."""
    s = (text or "").strip().lower()
    for a, b in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss")):
        s = s.replace(a, b)
    return re.sub(r"\s+", " ", s.replace(".", " ")).strip()


def _lit(v):
    return "'" + str(v).replace("'", "''") + "'"


def _fname(layer, wanted):
    """Feldnamen case-tolerant ermitteln (ogr2ogr/PostGIS)."""
    if layer is None:
        return None
    for f in layer.fields():
        if f.name().lower() == wanted.lower():
            return f.name()
    return None


def _numeric(layer, fname):
    try:
        return layer.fields().at(layer.fields().indexFromName(fname)).isNumeric()
    except Exception:
        return False


def _req(expr, limit=0):
    r = QgsFeatureRequest().setFilterExpression(expr)
    try:
        r.setFlags(QgsFeatureRequest.NoGeometry)
    except Exception:
        pass
    if limit:
        r.setLimit(limit)
    return r


# ── Layer beschaffen ────────────────────────────────────────────────────────

def layers_from_project():
    """Die drei benötigten Layer aus dem geladenen Projekt."""
    def one(name):
        ls = QgsProject.instance().mapLayersByName(name)
        return ls[0] if ls else None
    return one(REF_LAYER), one(SYN_LAYER), one(AG_LAYER)


def layers_from_gpkg(path):
    """Die drei Layer direkt aus einer GeoPackage-Datei (z. B. Plugin-Vorlage)."""
    def one(name):
        lyr = QgsVectorLayer(f"{path}|layername={name}", name, "ogr")
        return lyr if lyr.isValid() else None
    return one(REF_LAYER), one(SYN_LAYER), one(AG_LAYER)


# ── Auflösung / Validierung ─────────────────────────────────────────────────

def valid_groups(ag_layer):
    """Menge gültiger Artengruppen-Schlüssel (listitemid)."""
    out = set()
    if ag_layer is None:
        return out
    f = _fname(ag_layer, "listitemid")
    if not f:
        return out
    for feat in ag_layer.getFeatures():
        v = feat[f]
        if v not in (None, ""):
            out.add(str(v))
    return out


def group_names(ag_layer):
    res = {}
    if ag_layer is None:
        return res
    li, tm = _fname(ag_layer, "listitemid"), _fname(ag_layer, "term")
    if not (li and tm):
        return res
    for feat in ag_layer.getFeatures():
        res[str(feat[li])] = feat[tm]
    return res


def find_species(ref_layer, ag_layer, text, limit=30):
    """
    Kandidaten für eine Zielart suchen.

    Rückgabe: Liste von dicts mit entityid, anzeigename, parentid (GÜLTIG),
    artengruppe. Arten ohne gültige Artengruppe werden ausgeschlossen – sie
    wären im Formular ohnehin nicht erreichbar.
    """
    if ref_layer is None or len((text or "").strip()) < 2:
        return []
    key = _fname(ref_layer, "entityid")
    disp = _fname(ref_layer, "anzeigename")
    par = _fname(ref_layer, "parentid")
    fields = [x for x in (_fname(ref_layer, n)
                          for n in ("term", "Name_deutsch", "anzeigename")) if x]
    if not (key and disp and par):
        return []
    needle = str(text).strip().replace("'", "''")
    expr = " OR ".join('"%s" ILIKE %s' % (f, "'%" + needle + "%'") for f in fields)

    ok_groups = valid_groups(ag_layer)
    names = group_names(ag_layer)
    seen, out = set(), []
    for feat in ref_layer.getFeatures(_req(expr)):
        eid, pid = feat[key], str(feat[par])
        if eid in seen or pid not in ok_groups:
            continue
        seen.add(eid)
        out.append({"entityid": eid, "anzeigename": feat[disp],
                    "parentid": pid, "artengruppe": names.get(pid, pid)})
        if len(out) >= limit:
            break
    out.sort(key=lambda d: str(d["anzeigename"]).lower())
    return out


def resolve_exact(ref_layer, ag_layer, text):
    """
    Zielart exakt auflösen (für den Dateiimport).

    Rückgabe: (treffer, kandidaten). 'treffer' ist gesetzt, wenn genau eine
    Art eindeutig passt – sonst None und 'kandidaten' zur Anzeige.
    """
    if ref_layer is None:
        return None, []
    key = _fname(ref_layer, "entityid")
    # 1) direkte entityid?
    t = str(text).strip()
    if t.isdigit():
        for feat in ref_layer.getFeatures(_req('"%s" = %s' % (key, int(t)), 1)):
            cand = find_species(ref_layer, ag_layer, str(feat[_fname(ref_layer, "term")]))
            for c in cand:
                if str(c["entityid"]) == t:
                    return c, cand
    # 2) exakter Namensvergleich (normalisiert)
    cands = find_species(ref_layer, ag_layer, t, limit=50)
    n = normalize(t)
    exact = []
    for c in cands:
        for feat in ref_layer.getFeatures(
                _req('"%s" = %s' % (key, c["entityid"] if _numeric(ref_layer, key)
                                    else _lit(c["entityid"])), 1)):
            for fld in ("term", "Name_deutsch", "anzeigename"):
                f = _fname(ref_layer, fld)
                if f and normalize(feat[f]) == n:
                    exact.append(c)
                    break
    uniq = {str(c["entityid"]): c for c in exact}
    if len(uniq) == 1:
        return list(uniq.values())[0], cands
    return None, cands


def existing_terms(syn_layer):
    """Vorhandene Suchbegriffe (normalisiert) -> Originalschreibweise."""
    res = {}
    if syn_layer is None:
        return res
    f = _fname(syn_layer, "synonym_term")
    if not f:
        return res
    for feat in syn_layer.getFeatures():
        res[normalize(feat[f])] = feat[f]
    return res


def current_user():
    """Bearbeitername: QGIS-Benutzer, sonst Anmeldename des Systems."""
    for fn in ("userFullName", "userLoginName"):
        try:
            v = getattr(QgsApplication, fn)()
            if v:
                return str(v)
        except Exception:
            pass
    try:
        import getpass
        return getpass.getuser()
    except Exception:
        return ""


def ensure_audit_columns(syn_layer):
    """
    Nachweisspalten ergänzen, falls die Tabelle noch aus einer älteren
    Version stammt. Rückgabe: Liste der neu angelegten Spalten.
    """
    if syn_layer is None:
        return []
    neu = [n for n in ("erfasst_von", "erfasst_am") if not _fname(syn_layer, n)]
    if not neu:
        return []
    try:
        ok = syn_layer.dataProvider().addAttributes(
            [QgsField(n, QVariant.String) for n in neu])
        syn_layer.updateFields()
        if ok:
            log("Nachweisspalten ergänzt: %s" % ", ".join(neu))
            return neu
    except Exception as e:
        log("Nachweisspalten konnten nicht ergänzt werden: %s" % e)
    return []


def add_synonyms(syn_layer, entries, quelle=None, bearbeiter=None):
    """
    Synonyme schreiben – inklusive Nachweis (wer/wann).

    entries: Liste von dicts mit 'term', 'entityid', 'parentid'.
    Rückgabe: (anzahl_geschrieben, fehlerliste)
    """
    if syn_layer is None:
        return 0, ["Layer 'Arten_Synonyme' nicht verfügbar."]
    ensure_audit_columns(syn_layer)
    f_term = _fname(syn_layer, "synonym_term")
    f_key = _fname(syn_layer, "akzeptiert_entityid")
    f_par = _fname(syn_layer, "parentid")
    f_art = _fname(syn_layer, "art")
    f_qu = _fname(syn_layer, "quelle")
    f_von = _fname(syn_layer, "erfasst_von")
    f_am = _fname(syn_layer, "erfasst_am")
    bearbeiter = (bearbeiter or current_user() or "").strip()
    stempel = date.today().strftime("%Y-%m-%d")
    if not (f_term and f_key and f_par):
        return 0, ["Tabelle 'Arten_Synonyme' hat nicht die erwarteten Spalten."]

    quelle = quelle or ("manuell %s" % date.today().strftime("%Y-%m-%d"))
    have = existing_terms(syn_layer)
    feats, errs = [], []
    for e in entries:
        term = (e.get("term") or "").strip()
        if not term:
            continue
        if normalize(term) in have:
            errs.append("„%s“ ist bereits vorhanden – übersprungen." % term)
            continue
        feat = QgsFeature(syn_layer.fields())
        feat[f_term] = term
        feat[f_key] = e["entityid"]
        feat[f_par] = str(e["parentid"])
        if f_art:
            feat[f_art] = e.get("art") or "Synonym"
        if f_qu:
            feat[f_qu] = e.get("quelle") or quelle
        if f_von:
            feat[f_von] = e.get("erfasst_von") or bearbeiter
        if f_am:
            feat[f_am] = e.get("erfasst_am") or stempel
        feats.append(feat)
        have[normalize(term)] = term

    if not feats:
        return 0, errs

    started = syn_layer.startEditing()
    ok = syn_layer.addFeatures(feats)
    if ok:
        if not syn_layer.commitChanges():
            errs.append("Speichern fehlgeschlagen: %s"
                        % "; ".join(syn_layer.commitErrors()))
            syn_layer.rollBack()
            return 0, errs
    else:
        if started:
            syn_layer.rollBack()
        errs.append("Datensätze konnten nicht angelegt werden.")
        return 0, errs
    log("%d Synonym(e) ergänzt (%s)" % (len(feats), quelle))
    return len(feats), errs


def delete_synonyms(syn_layer, terms):
    """Synonyme anhand ihres Suchbegriffs entfernen."""
    if syn_layer is None or not terms:
        return 0
    f = _fname(syn_layer, "synonym_term")
    want = {normalize(t) for t in terms}
    ids = [feat.id() for feat in syn_layer.getFeatures()
           if normalize(feat[f]) in want]
    if not ids:
        return 0
    syn_layer.startEditing()
    syn_layer.deleteFeatures(ids)
    if not syn_layer.commitChanges():
        syn_layer.rollBack()
        return 0
    return len(ids)


# ── Prüfung des Bestands ────────────────────────────────────────────────────

def check(syn_layer, ref_layer, ag_layer):
    """Bestand prüfen. Rückgabe: dict mit Befunden."""
    res = {"gesamt": 0, "ungueltige_gruppe": [], "zielart_fehlt": [],
           "doppelt": [], "gruppe_abweichend": []}
    if syn_layer is None or ref_layer is None:
        return res
    f_term = _fname(syn_layer, "synonym_term")
    f_key = _fname(syn_layer, "akzeptiert_entityid")
    f_par = _fname(syn_layer, "parentid")
    r_key = _fname(ref_layer, "entityid")
    r_par = _fname(ref_layer, "parentid")

    ok_groups = valid_groups(ag_layer)
    arten = {}
    for feat in ref_layer.getFeatures():
        arten.setdefault(str(feat[r_key]), set()).add(str(feat[r_par]))

    seen = {}
    for feat in syn_layer.getFeatures():
        res["gesamt"] += 1
        term, key, par = feat[f_term], str(feat[f_key]), str(feat[f_par])
        n = normalize(term)
        if n in seen:
            res["doppelt"].append(term)
        seen[n] = term
        if key not in arten:
            res["zielart_fehlt"].append(term)
            continue
        if par not in ok_groups:
            res["ungueltige_gruppe"].append(term)
        elif par not in arten[key]:
            res["gruppe_abweichend"].append(term)
    return res


def repair_groups(syn_layer, ref_layer, ag_layer):
    """
    Ungültige/abweichende parentid automatisch auf eine gültige Artengruppe
    der jeweiligen Zielart setzen. Rückgabe: (korrigiert, nicht_korrigierbar)
    """
    if syn_layer is None or ref_layer is None:
        return 0, []
    f_term = _fname(syn_layer, "synonym_term")
    f_key = _fname(syn_layer, "akzeptiert_entityid")
    f_par = _fname(syn_layer, "parentid")
    r_key = _fname(ref_layer, "entityid")
    r_par = _fname(ref_layer, "parentid")
    idx_par = syn_layer.fields().indexFromName(f_par)

    ok_groups = valid_groups(ag_layer)
    arten = {}
    for feat in ref_layer.getFeatures():
        arten.setdefault(str(feat[r_key]), set()).add(str(feat[r_par]))

    upd, bad = {}, []
    for feat in syn_layer.getFeatures():
        key, par = str(feat[f_key]), str(feat[f_par])
        groups = arten.get(key, set())
        if par in ok_groups and par in groups:
            continue
        valid = [g for g in groups if g in ok_groups]
        if valid:
            upd[feat.id()] = {idx_par: valid[0]}
        else:
            bad.append(feat[f_term])
    if upd:
        syn_layer.startEditing()
        syn_layer.dataProvider().changeAttributeValues(upd)
        if not syn_layer.commitChanges():
            syn_layer.rollBack()
            return 0, bad
    return len(upd), bad


# ── Datei-Import ────────────────────────────────────────────────────────────

def read_table(path):
    """
    CSV oder XLSX einlesen. Erwartet zwei Spalten: Suchbegriff und Zielart
    (wissenschaftlicher/deutscher Name oder entityid). Kopfzeile optional.
    Rückgabe: Liste von (suchbegriff, zielart).
    """
    ext = os.path.splitext(path)[1].lower()
    rows = []
    if ext in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise RuntimeError("Für XLSX wird 'openpyxl' benötigt – "
                               "bitte CSV verwenden.")
        ws = load_workbook(path, read_only=True, data_only=True).active
        for r in ws.iter_rows(values_only=True):
            if r and any(x not in (None, "") for x in r[:2]):
                rows.append([("" if r[0] is None else str(r[0])),
                             ("" if len(r) < 2 or r[1] is None else str(r[1]))])
    else:
        with open(path, newline="", encoding="utf-8-sig") as fh:
            sample = fh.read(2048)
            fh.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
            except Exception:
                dialect = csv.excel
                dialect.delimiter = ";"
            for r in csv.reader(fh, dialect):
                if r and any(x.strip() for x in r[:2]):
                    rows.append([r[0], r[1] if len(r) > 1 else ""])
    # Kopfzeile erkennen
    if rows and normalize(rows[0][0]) in ("synonym", "synonym term", "suchbegriff",
                                          "synonym_term", "begriff"):
        rows = rows[1:]
    return [(a.strip(), b.strip()) for a, b in rows]


def prepare_import(rows, ref_layer, syn_layer, ag_layer):
    """
    Import vorbereiten (Trockenlauf).
    Rückgabe: (bereit, probleme) – 'bereit' sind fertige Einträge für
    add_synonyms(), 'probleme' eine Liste (suchbegriff, grund, kandidaten).
    """
    have = existing_terms(syn_layer)
    ready, problems = [], []
    for term, target in rows:
        if not term:
            continue
        if normalize(term) in have:
            problems.append((term, "bereits vorhanden", []))
            continue
        if not target:
            problems.append((term, "keine Zielart angegeben", []))
            continue
        hit, cands = resolve_exact(ref_layer, ag_layer, target)
        if hit is None:
            grund = ("Zielart nicht gefunden" if not cands
                     else "Zielart nicht eindeutig (%d Kandidaten)" % len(cands))
            problems.append((term, grund, [c["anzeigename"] for c in cands[:5]]))
            continue
        ready.append({"term": term, "entityid": hit["entityid"],
                      "parentid": hit["parentid"],
                      "ziel": hit["anzeigename"], "gruppe": hit["artengruppe"]})
        have[normalize(term)] = term
    return ready, problems


# ── Abgleich zwischen zwei Beständen ────────────────────────────────────────

def read_all(syn_layer):
    """Bestand als dict: normalisierter Begriff -> Datensatz."""
    out = {}
    if syn_layer is None:
        return out
    f_term = _fname(syn_layer, "synonym_term")
    f_key = _fname(syn_layer, "akzeptiert_entityid")
    f_par = _fname(syn_layer, "parentid")
    f_art = _fname(syn_layer, "art")
    f_qu = _fname(syn_layer, "quelle")
    f_von = _fname(syn_layer, "erfasst_von")
    f_am = _fname(syn_layer, "erfasst_am")
    for feat in syn_layer.getFeatures():
        term = feat[f_term]
        if term in (None, ""):
            continue
        out[normalize(term)] = {
            "term": term,
            "entityid": feat[f_key],
            "parentid": str(feat[f_par]),
            "art": feat[f_art] if f_art else "",
            "quelle": feat[f_qu] if f_qu else "",
            "erfasst_von": (feat[f_von] if f_von else "") or "",
            "erfasst_am": (feat[f_am] if f_am else "") or "",
        }
    return out


def species_names(ref_layer, entityids):
    """entityid -> akzeptierter Anzeigename (für die Abgleichsanzeige)."""
    res = {}
    if ref_layer is None or not entityids:
        return res
    key = _fname(ref_layer, "entityid")
    disp = _fname(ref_layer, "anzeigename")
    if not (key and disp):
        return res
    num = _numeric(ref_layer, key)
    vals = ", ".join((str(int(e)) if num else _lit(e)) for e in entityids)
    for feat in ref_layer.getFeatures(_req('"%s" IN (%s)' % (key, vals))):
        res.setdefault(str(feat[key]), feat[disp])
    return res


def compare(syn_here, syn_there, ref_layer=None):
    """
    Zwei Synonymbestände vergleichen.

    Rückgabe: dict mit
      'nur_hier'   – Einträge, die nur im Arbeitsbestand stehen
      'nur_dort'   – Einträge, die nur im Vergleichsbestand stehen
      'abweichend' – gleicher Suchbegriff, andere Zielart
    """
    a, b = read_all(syn_here), read_all(syn_there)
    nur_hier = [a[k] for k in a if k not in b]
    nur_dort = [b[k] for k in b if k not in a]
    abweichend = []
    for k in a:
        if k in b and str(a[k]["entityid"]) != str(b[k]["entityid"]):
            abweichend.append({"term": a[k]["term"],
                               "hier": a[k], "dort": b[k]})
    if ref_layer is not None:
        ids = {str(x["entityid"]) for x in nur_hier + nur_dort}
        ids |= {str(x["hier"]["entityid"]) for x in abweichend}
        ids |= {str(x["dort"]["entityid"]) for x in abweichend}
        names = species_names(ref_layer, [i for i in ids if i not in (None, "None", "")])
        for x in nur_hier + nur_dort:
            x["ziel"] = names.get(str(x["entityid"]), str(x["entityid"]))
        for x in abweichend:
            x["ziel_hier"] = names.get(str(x["hier"]["entityid"]), "")
            x["ziel_dort"] = names.get(str(x["dort"]["entityid"]), "")
    for x in nur_hier + nur_dort:
        x.setdefault("ziel", str(x["entityid"]))
    return {"nur_hier": sorted(nur_hier, key=lambda d: str(d["term"]).lower()),
            "nur_dort": sorted(nur_dort, key=lambda d: str(d["term"]).lower()),
            "abweichend": sorted(abweichend, key=lambda d: str(d["term"]).lower())}


def transfer(entries, target_layer):
    """
    Einträge in einen anderen Bestand übertragen – Nachweis bleibt erhalten
    (erfasst_von/erfasst_am der Quelle werden mitgenommen).
    """
    payload = []
    for e in entries:
        payload.append({
            "term": e["term"], "entityid": e["entityid"], "parentid": e["parentid"],
            "art": e.get("art") or "Synonym",
            "quelle": e.get("quelle") or "",
            "erfasst_von": e.get("erfasst_von") or "",
            "erfasst_am": e.get("erfasst_am") or "",
        })
    return add_synonyms(target_layer, payload)
